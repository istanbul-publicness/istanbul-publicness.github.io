"""
Compile publicness_kurallar.xlsx into rules.json, the file the computation reads.

The workbook is where people decide; rules.json is what the code runs on. The
split keeps the computation free of a spreadsheet library and makes every run
traceable: rules.json carries the workbook's SHA-256 and the time it was
compiled, and the computation writes both into its output.

Effective value of any row = the reviewers' KARAR if written, else the proposal.

Run with QGIS's Python (it has openpyxl), after every change to the workbook:
    "C:\\Program Files\\QGIS 3.44.12\\bin\\python-qgis-ltr.bat" compile_rules.py
"""

import datetime
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
BOOK = HERE / "publicness_kurallar.xlsx"
OUT = HERE / "rules.json"

CODE = {"açık-kamusal": "open_public", "yarı-kamusal": "quasi_public", "biletli": "ticketed",
        "açık-özel": "open_private", "davetli": "invitation", "erişimsiz": "inaccessible",
        "sınıflanmadı": None, "zemin dışı": "exclude", "taşıt yolu": "car",
        "doğal alan (skor dışı)": "natural"}
REGIMES = ["open_public", "quasi_public", "ticketed", "open_private", "invitation", "inaccessible"]


def rows(ws):
    data = list(ws.iter_rows(values_only=True))
    head = [str(h) if h is not None else "" for h in data[0]]
    return [dict(zip(head, r)) for r in data[1:] if any(v is not None for v in r)]


def effective(r, proposal_col):
    k = r.get("KARAR (siz)")
    return k if k not in (None, "") else r.get(proposal_col)


def code(label, where):
    if label in (None, ""):
        return None
    label = str(label).strip()
    if label not in CODE:
        raise ValueError(f"{where}: tanınmayan rejim '{label}'")
    return CODE[label]


def number(text, default):
    m = re.search(r"\d+(?:[.,]\d+)?", str(text or ""))
    return float(m.group(0).replace(",", ".")) if m else default


def main():
    sha = hashlib.sha256(BOOK.read_bytes()).hexdigest()
    wb = load_workbook(BOOK, data_only=True, read_only=True)

    # scores: the column the reviewers fill, plus the schemes kept for sensitivity
    pr = rows(wb["puanlar"])[:6]
    scores = {code(r["Rejim"], "puanlar"): float(r["KULLANILACAK PUAN (siz)"]) for r in pr}
    sens = {
        "yedide_bir": {code(r["Rejim"], "puanlar"): float(r["mevcut (yedide bir)"]) for r in pr},
        "alpha_0.7": {code(r["Rejim"], "puanlar"): float(r["erişim öncelikli (α=0,7)"]) for r in pr},
        "alpha_0.3": {code(r["Rejim"], "puanlar"): float(r["kontrol öncelikli (α=0,3)"]) for r in pr},
    }
    assert set(scores) == set(REGIMES), scores

    decisions = {r["No"]: r.get("KARAR (siz)") for r in rows(wb["kararlar"]) if r.get("No")}

    osm = {}
    for r in rows(wb["osm_kurallari"]):
        tag = str(r.get("OSM etiketi") or "").strip()
        if re.fullmatch(r"[a-z_:]+=[a-z_]+", tag):
            osm[tag] = code(effective(r, "Rejim"), f"osm {tag}")

    overture = {}
    for r in rows(wb["overture_sozluk"]):
        cat = r.get("Overture temel kategori")
        if cat:
            overture[cat] = code(effective(r, "Önerilen rejim"), f"overture {cat}")

    foursquare = {}
    for r in rows(wb["foursquare_sozluk"]):
        cat = r.get("Foursquare kategori")
        if cat:
            foursquare[cat] = code(effective(r, "Önerilen rejim"), f"foursquare {cat}")

    clean = {"drop_closed": True, "overture_min_conf": 0.5, "stale_before": "2019-01-01",
             "match_m": 50.0, "name_sim": 0.8, "non_ground_fsq": [], "non_ground_overture": []}
    for r in rows(wb["poi_temizlik"]):
        rule, src = str(r.get("Kural") or ""), r.get("Kaynak")
        val = effective(r, "Önerim")
        on = str(val).strip().lower() in ("evet", "at")
        if rule.startswith("Kapanmış"):
            clean["drop_closed"] = on
        elif rule.startswith("Overture güven"):
            clean["overture_min_conf"] = number(r.get("Değer"), 0.5) if on else 0.0
        elif rule.startswith("Eskime"):
            clean["stale_before"] = "2019-01-01" if on else None
        elif rule.startswith("Eşleştirme mesafesi"):
            clean["match_m"] = number(r.get("Değer"), 50.0)
        elif rule.startswith("İsim benzerliği"):
            clean["name_sim"] = number(r.get("Değer"), 0.8)
        elif rule.startswith("Zemin dışı kategori") and on:
            clean["non_ground_fsq" if src == "Foursquare" else "non_ground_overture"].append(r.get("Değer"))

    # who controls a named site, where OSM does not say (D07)
    control = {}
    if "kontrol" in wb.sheetnames:
        for r in rows(wb["kontrol"]):
            name, c = r.get("Mekân adı (OSM name)"), str(r.get("Kontrol") or "").strip().lower()
            if name and c in ("kamu", "özel"):
                control[str(name).strip()] = "public" if c == "kamu" else "private"

    # D19 review: a regime for a matched site, or 'uygulama' to leave it as it was
    d19 = {}
    if "d19_adaylar" in wb.sheetnames:
        for r in rows(wb["d19_adaylar"]):
            k = str(r.get("KARAR (siz)") or "").strip()
            if r.get("Alan") and k:
                d19[str(r["Alan"]).strip()] = "off" if k.lower() == "uygulama" else code(k, f"d19 {r['Alan']}")

    # hand corrections: an OSM element (w123 / r456) or a name, and the regime it really has
    fixes = {}
    if "duzeltmeler" in wb.sheetnames:
        for r in rows(wb["duzeltmeler"]):
            key = str(r.get("OSM no ya da ad") or "").strip()
            if key and r.get("Rejim"):
                fixes[key] = code(r["Rejim"], f"düzeltme {key}")

    out = {"source": BOOK.name, "sha256": sha,
           "compiled": datetime.datetime.now().isoformat(timespec="seconds"),
           "scores": scores, "sensitivity": sens, "decisions": decisions, "osm": osm,
           "overture": overture, "foursquare": foursquare, "cleaning": clean, "control": control,
           "d19": d19, "duzeltme": fixes}
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")

    def tally(d):
        c = {}
        for v in d.values():
            c[v or "sınıflanmadı"] = c.get(v or "sınıflanmadı", 0) + 1
        return ", ".join(f"{k} {n}" for k, n in sorted(c.items(), key=lambda kv: -kv[1]))

    print(f"yazıldı: {OUT.name} · kaynak {BOOK.name} · sha256 {sha[:12]}")
    print(f"puanlar: " + ", ".join(f"{k} {v:.3f}" for k, v in scores.items()))
    print(f"kararlar: {sum(1 for v in decisions.values() if v)}/{len(decisions)} dolu")
    print(f"D19 gözden geçirme: {len(d19)} karar · elle düzeltme: {len(fixes)}")
    print(f"kontrol: {len(control)} mekân ({', '.join(f'{k} → {v}' for k, v in control.items())})")
    print(f"OSM kuralı {len(osm)} · Overture {len(overture)} ({tally(overture)})")
    print(f"Foursquare {len(foursquare)} ({tally(foursquare)})")
    print(f"temizlik: {clean['non_ground_fsq'].__len__()} Foursquare + {len(clean['non_ground_overture'])} Overture "
          f"zemin dışı kategori · güven ≥ {clean['overture_min_conf']} · eskime {clean['stale_before']}")


if __name__ == "__main__":
    main()
